from pathlib import Path
from docx import Document
from ..models import ContentUnit, ParsedDocument

class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(str(path))
        units, current_heading = [], None
        for i, para in enumerate(doc.paragraphs, 1):
            text = " ".join((para.text or "").split())
            if not text:
                continue
            style = getattr(para.style, "name", "") or ""
            if style.lower().startswith("heading") or "title" in style.lower():
                current_heading = text
            units.append(ContentUnit(text, "paragraph",
                                     {"paragraph_number": i, "style": style}, current_heading))
        for tno, table in enumerate(doc.tables, 1):
            for rno, row in enumerate(table.rows, 1):
                cells = []
                for cno, cell in enumerate(row.cells, 1):
                    value = " ".join((cell.text or "").split())
                    if value:
                        cells.append(f"Cell {cno}: {value}")
                if cells:
                    units.append(ContentUnit(f"Table {tno}, row {rno}: " + " | ".join(cells),
                                             "word_table_row",
                                             {"table_number": tno, "row_number": rno},
                                             current_heading))
        return ParsedDocument(units)
