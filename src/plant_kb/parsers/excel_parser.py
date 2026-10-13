from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from ..models import ContentUnit, ParsedDocument

KEYWORDS = ("hazard", "risk", "activity", "parameter", "equipment", "description",
            "likelihood", "severity", "control", "department", "serial", "date",
            "process", "task", "recommendation", "action", "s.no")

def clean(v):
    return " ".join(str(v).replace("\n", " ").split()) if v is not None else ""

def detect_header(preview):
    candidates = []
    for row_no, row in enumerate(preview, 1):
        vals = [v for v in row if v]
        if len(vals) < 3:
            continue
        joined = " ".join(vals).lower()
        score = len(vals) + sum(2 for kw in KEYWORDS if kw in joined)
        candidates.append((score, row_no))
    return max(candidates)[1] if candidates else None

class ExcelParser:
    def parse(self, path: Path) -> ParsedDocument:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        units, warnings = [], []
        try:
            for ws in wb.worksheets:
                preview = [[clean(v) for v in row]
                           for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 30), values_only=True)]
                header_no = detect_header(preview)
                header = preview[header_no - 1] if header_no and header_no <= len(preview) else None
                nonempty_count = 0
                for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
                    values = [clean(v) for v in row]
                    populated = [(i, v) for i, v in enumerate(values, 1) if v]
                    if not populated:
                        continue
                    nonempty_count += 1
                    fields = []
                    for col_no, value in populated:
                        if header and row_no > header_no and col_no <= len(header) and header[col_no - 1]:
                            label = header[col_no - 1]
                        else:
                            label = f"Column {get_column_letter(col_no)}"
                        fields.append(f"{label}: {value}")
                    text = f"Sheet '{ws.title}', row {row_no}: " + " | ".join(fields)
                    units.append(ContentUnit(text, "spreadsheet_row",
                        {"sheet_name": ws.title, "row_number": row_no, "detected_header_row": header_no},
                        ws.title))
                if not nonempty_count:
                    warnings.append(f"Sheet '{ws.title}' is empty.")
                elif not header_no:
                    warnings.append(f"No likely header row detected in sheet '{ws.title}'; column letters were preserved.")
                elif header_no > 15:
                    warnings.append(f"Header row for sheet '{ws.title}' was heuristically detected at row {header_no}; validate.")
        finally:
            wb.close()
        return ParsedDocument(units, warnings)
